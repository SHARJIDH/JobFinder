"""ATS-friendly rendering: single column, standard headings, no tables/images."""
from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document as Docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer


def _contact_line(c: dict) -> str:
    keys = ("email", "phone", "location", "linkedin", "github", "portfolio")
    return " | ".join(c.get(k, "") for k in keys if c.get(k))


def _dates(start: str, end: str) -> str:
    if not start and not end:
        return ""
    return f"{start or '?'} – {end or 'Present'}"


# ------------------------------------------------------------------ DOCX

def _docx_base() -> Docx:
    d = Docx()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for s in d.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Pt(46)
    return d


def _heading(d: Docx, text: str) -> None:
    p = d.add_paragraph()
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(11.5)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)


def resume_docx(content: dict, path: Path) -> None:
    d = _docx_base()
    c = content.get("contact", {})
    name = d.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = name.add_run(c.get("full_name") or "Your Name")
    r.bold, r.font.size = True, Pt(16)
    line = _contact_line(c)
    if line:
        p = d.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if content.get("summary"):
        _heading(d, "Summary")
        d.add_paragraph(content["summary"])
    if content.get("skills"):
        _heading(d, "Skills")
        d.add_paragraph(", ".join(content["skills"]))
    if content.get("experience"):
        _heading(d, "Experience")
        for x in content["experience"]:
            p = d.add_paragraph()
            p.add_run(f"{x['title']} — {x['company']}").bold = True
            meta = ", ".join(v for v in (x.get("location"), _dates(x.get("start", ""), x.get("end", ""))) if v)
            if meta:
                p.add_run(f"  ({meta})")
            for b in x.get("bullets", []):
                d.add_paragraph(b["text"], style="List Bullet")
    if content.get("projects"):
        _heading(d, "Projects")
        for pr in content["projects"]:
            p = d.add_paragraph()
            p.add_run(pr["name"]).bold = True
            if pr.get("technologies"):
                p.add_run(f"  ({', '.join(pr['technologies'])})")
            if pr.get("url"):
                p.add_run(f"  {pr['url']}")
            for b in pr.get("bullets", []):
                d.add_paragraph(b["text"], style="List Bullet")
    if content.get("education"):
        _heading(d, "Education")
        for e in content["education"]:
            p = d.add_paragraph()
            p.add_run(f"{e['degree']} — {e['institution']}").bold = True
            extra = ", ".join(v for v in (_dates(e.get("start", ""), e.get("end", "")), e.get("grade", "")) if v)
            if extra:
                p.add_run(f"  ({extra})")
    if content.get("certifications"):
        _heading(d, "Certifications")
        for cert in content["certifications"]:
            d.add_paragraph(" — ".join(v for v in (cert.get("name"), cert.get("issuer"), cert.get("date")) if v), style="List Bullet")
    if content.get("achievements"):
        _heading(d, "Achievements")
        for a in content["achievements"]:
            d.add_paragraph(a, style="List Bullet")
    d.save(str(path))


def cover_docx(letter: dict, path: Path) -> None:
    d = _docx_base()
    c = letter.get("contact", {})
    p = d.add_paragraph()
    p.add_run(c.get("full_name") or "Your Name").bold = True
    if _contact_line(c):
        d.add_paragraph(_contact_line(c))
    d.add_paragraph(f"Re: {letter['title']} — {letter['company']}")
    d.add_paragraph("Dear Hiring Team,")
    for para in letter["paragraphs"]:
        d.add_paragraph(para)
    d.add_paragraph("Sincerely,")
    d.add_paragraph(c.get("full_name") or "")
    d.save(str(path))


# ------------------------------------------------------------------ PDF

def _styles():
    ss = getSampleStyleSheet()
    return {
        "name": ParagraphStyle("name", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=16, spaceAfter=2, alignment=TA_CENTER),
        "contact": ParagraphStyle("contact", parent=ss["Normal"], fontName="Helvetica", fontSize=9, alignment=TA_CENTER, spaceAfter=6),
        "h": ParagraphStyle("h", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=11, spaceBefore=8, spaceAfter=3),
        "body": ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.8, leading=12.5),
        "role": ParagraphStyle("role", parent=ss["Normal"], fontName="Helvetica", fontSize=10, leading=13, spaceBefore=4),
    }


def _bullets(items: list[str], st) -> ListFlowable:
    return ListFlowable([ListItem(Paragraph(escape(t), st["body"]), leftIndent=10) for t in items],
                        bulletType="bullet", start="•", leftIndent=12, bulletFontSize=8)


def _doc(path: Path) -> SimpleDocTemplate:
    return SimpleDocTemplate(str(path), pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm,
                             topMargin=14 * mm, bottomMargin=14 * mm, title="Resume", author="")


def resume_pdf(content: dict, path: Path) -> None:
    st = _styles()
    c = content.get("contact", {})
    flow = [Paragraph(escape(c.get("full_name") or "Your Name"), st["name"])]
    if _contact_line(c):
        flow.append(Paragraph(escape(_contact_line(c)), st["contact"]))
    if content.get("summary"):
        flow += [Paragraph("SUMMARY", st["h"]), Paragraph(escape(content["summary"]), st["body"])]
    if content.get("skills"):
        flow += [Paragraph("SKILLS", st["h"]), Paragraph(escape(", ".join(content["skills"])), st["body"])]
    if content.get("experience"):
        flow.append(Paragraph("EXPERIENCE", st["h"]))
        for x in content["experience"]:
            meta = ", ".join(v for v in (x.get("location"), _dates(x.get("start", ""), x.get("end", ""))) if v)
            flow.append(Paragraph(f"<b>{escape(x['title'])} — {escape(x['company'])}</b>" + (f"  ({escape(meta)})" if meta else ""), st["role"]))
            if x.get("bullets"):
                flow.append(_bullets([b["text"] for b in x["bullets"]], st))
    if content.get("projects"):
        flow.append(Paragraph("PROJECTS", st["h"]))
        for p in content["projects"]:
            extra = f"  ({escape(', '.join(p['technologies']))})" if p.get("technologies") else ""
            flow.append(Paragraph(f"<b>{escape(p['name'])}</b>{extra}" + (f"  {escape(p['url'])}" if p.get("url") else ""), st["role"]))
            if p.get("bullets"):
                flow.append(_bullets([b["text"] for b in p["bullets"]], st))
    if content.get("education"):
        flow.append(Paragraph("EDUCATION", st["h"]))
        for e in content["education"]:
            extra = ", ".join(v for v in (_dates(e.get("start", ""), e.get("end", "")), e.get("grade", "")) if v)
            flow.append(Paragraph(f"<b>{escape(e['degree'])} — {escape(e['institution'])}</b>" + (f"  ({escape(extra)})" if extra else ""), st["role"]))
    if content.get("certifications"):
        flow += [Paragraph("CERTIFICATIONS", st["h"]),
                 _bullets([" — ".join(v for v in (x.get("name"), x.get("issuer"), x.get("date")) if v) for x in content["certifications"]], st)]
    if content.get("achievements"):
        flow += [Paragraph("ACHIEVEMENTS", st["h"]), _bullets(content["achievements"], st)]
    _doc(path).build(flow)


def cover_pdf(letter: dict, path: Path) -> None:
    st = _styles()
    c = letter.get("contact", {})
    flow = [Paragraph(f"<b>{escape(c.get('full_name') or 'Your Name')}</b>", st["body"])]
    if _contact_line(c):
        flow.append(Paragraph(escape(_contact_line(c)), st["body"]))
    flow += [Spacer(1, 10), Paragraph(escape(f"Re: {letter['title']} — {letter['company']}"), st["body"]), Spacer(1, 8),
             Paragraph("Dear Hiring Team,", st["body"]), Spacer(1, 6)]
    for para in letter["paragraphs"]:
        flow += [Paragraph(escape(para), st["body"]), Spacer(1, 6)]
    flow += [Paragraph("Sincerely,", st["body"]), Paragraph(escape(c.get("full_name") or ""), st["body"])]
    _doc(path).build(flow)
